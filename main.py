import discord
from discord.ext import commands, tasks
from discord import app_commands
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import io
import random
import os
import json
from datetime import datetime, timedelta

# ================== CONFIGURACIÓN ==================
TOKEN = os.getenv("DISCORD_TOKEN")

CANAL_CHISTES   = 1549214401830322186
CANAL_POEMAS    = 1544141490152931479
CANAL_ARTE      = 1437938654747164855
CANAL_INFO      = 1544140921061380116
CANAL_NIVELES   = 1549209849034973244
# ===================================================

intents = discord.Intents.default()
intents.message_content = True
intents.members = True
intents.presences = True

bot = commands.Bot(command_prefix="!", intents=intents)

# ========== DATOS ==========
DATA_FILE = "datos.json"

def cargar_datos():
    try:
        with open(DATA_FILE, "r") as f:
            return json.load(f)
    except:
        return {"niveles": {}, "mensajes_totales": 0, "actividad": {}}

def guardar_datos(data):
    with open(DATA_FILE, "w") as f:
        json.dump(data, f, indent=4)

def calcular_nivel(xp):
    return int((xp ** 0.5) / 2)

cooldown_xp = {}

# ========== LISTAS ==========
chistes = [
    "¡Hola a todos! 👋 ¿Cómo van?",
    "¡Buenas! Espero que estén teniendo un gran día 😊",
    "Saludos desde el bot 🤖",
    "Chiste: ¿Qué le dice un jardinero a otro? ¡Estamos en las mismas!",
    "Chiste: ¿Por qué los pájaros no usan Facebook? Porque ya tienen Twitter 🐦",
    "Chiste: ¿Qué hace una abeja en el gimnasio? ¡Zum-ba!",
    "Chiste: ¿Cuál es el colmo de un electricista? No encontrar su corriente.",
    "Chiste: ¿Por qué el libro de matemáticas está triste? Porque tiene muchos problemas.",
    "¡Hola comunidad! Recuerden hidratarse 💧",
    "Chiste: ¿Qué le dijo una impresora a otra? ¿Esa hoja es tuya o es una impresión mía?",
    "Chiste: ¿Por qué las focas miran siempre hacia arriba? ¡Porque ahí están los focos!",
]

poemas = [
    "En la quietud de la noche,\nuna estrella me guió.\nY en su luz encontré\nel camino que busqué.",
    "El viento susurra secretos\nentre las hojas del tiempo.\nEscucha con el corazón\ny hallarás lo que anhelas.",
    "No temas a la lluvia,\nella limpia el alma.\nDespués del diluvio\nsiempre nace un nuevo día.",
    "Caminé entre sombras\ny encontré mi propia luz.\nA veces hay que perderse\npara poder encontrarse.",
    "Las palabras son puentes\nque unen almas distantes.\nHabla con ternura\ny el mundo se hará más grande.",
    "En cada final hay un comienzo,\nen cada lágrima una semilla.\nConfía en el ciclo de la vida,\ntodo vuelve a florecer.",
    "El silencio también habla,\nsolo hay que aprender a oírlo.\nEn su profundidad se esconden\nlas verdades más sinceras.",
]

calificaciones_arte = [
    "¡Qué hermoso! Esto tiene vibes de Emilia ❄️ (9/10)",
    "Rem estaría orgullosa de este arte 💙 (10/10)",
    "¡Increíble! Parece sacado directamente de Re:Zero (9.5/10)",
    "Esto grita Subaru en modo determinado 🔥 (8.5/10)",
    "Arte de calidad isekai ✨ (9/10)",
    "Beatrice aprobaría esto completamente 📚 (10/10)",
    "Muy bueno, pero le faltó un poco de drama de Return by Death (8/10)",
    "¡Espectacular! Esto merece estar en el opening (9.8/10)",
    "Calidad premium de otro mundo 🌌 (9/10)",
    "Ram diría que no está mal... y eso es mucho 😳 (8.7/10)",
]

# ========== EVENTOS ==========
@bot.event
async def on_message(message):
    if message.author.bot:
        return

    data = cargar_datos()
    data["mensajes_totales"] = data.get("mensajes_totales", 0) + 1
    guardar_datos(data)

    # --- Sistema de niveles ---
    user_id = str(message.author.id)
    ahora = datetime.utcnow()

    if user_id not in cooldown_xp or ahora - cooldown_xp[user_id] >= timedelta(seconds=60):
        cooldown_xp[user_id] = ahora
        if "niveles" not in data:
            data["niveles"] = {}
        if user_id not in data["niveles"]:
            data["niveles"][user_id] = {"xp": 0, "nivel": 0, "nombre": str(message.author)}

        data["niveles"][user_id]["xp"] += random.randint(15, 25)
        data["niveles"][user_id]["nombre"] = str(message.author)

        nivel_anterior = data["niveles"][user_id]["nivel"]
        nuevo_nivel = calcular_nivel(data["niveles"][user_id]["xp"])
        data["niveles"][user_id]["nivel"] = nuevo_nivel
        guardar_datos(data)

        if nuevo_nivel > nivel_anterior:
            canal = bot.get_channel(CANAL_NIVELES)
            if canal:
                await canal.send(f"🎉 ¡Felicidades {message.author.mention}! Has subido al **nivel {nuevo_nivel}**!")

    # --- Calificar arte ---
    if message.channel.id == CANAL_ARTE and message.attachments:
        for att in message.attachments:
            if att.content_type and att.content_type.startswith("image/"):
                await message.reply(random.choice(calificaciones_arte))
                break

    await bot.process_commands(message)

# ========== TAREAS ==========
@tasks.loop(minutes=10)
async def enviar_chistes():
    canal = bot.get_channel(CANAL_CHISTES)
    if canal:
        await canal.send(random.choice(chistes))

@tasks.loop(minutes=15)
async def enviar_poemas():
    canal = bot.get_channel(CANAL_POEMAS)
    if canal:
        await canal.send(f"📜 **Poema del momento:**\n\n{random.choice(poemas)}")

@tasks.loop(minutes=30)
async def enviar_info():
    canal = bot.get_channel(CANAL_INFO)
    if not canal or not canal.guild:
        return

    guild = canal.guild
    online = sum(1 for m in guild.members if m.status != discord.Status.offline)
    data = cargar_datos()

    embed = discord.Embed(title="📊 Estado del Servidor", color=discord.Color.green(), timestamp=datetime.utcnow())
    embed.add_field(name="👥 Miembros totales", value=str(guild.member_count), inline=True)
    embed.add_field(name="🟢 En línea", value=str(online), inline=True)
    embed.add_field(name="💬 Mensajes totales", value=str(data.get("mensajes_totales", 0)), inline=False)
    embed.set_footer(text="Actualización cada 30 minutos")
    await canal.send(embed=embed)

# ========== COMANDOS ==========
@bot.tree.command(name="nivel", description="Muestra tu nivel y XP")
async def nivel(interaction: discord.Interaction):
    data = cargar_datos()
    user_id = str(interaction.user.id)
    niveles = data.get("niveles", {})

    if user_id not in niveles:
        await interaction.response.send_message("Todavía no tienes XP. ¡Escribe mensajes para ganar niveles!", ephemeral=True)
        return

    info = niveles[user_id]
    embed = discord.Embed(title=f"📊 Nivel de {interaction.user.display_name}", color=discord.Color.blue())
    embed.add_field(name="Nivel", value=f"**{info['nivel']}**", inline=True)
    embed.add_field(name="XP", value=f"**{info['xp']}**", inline=True)
    embed.set_thumbnail(url=interaction.user.display_avatar.url)
    await interaction.response.send_message(embed=embed)

@bot.tree.command(name="ranking", description="Top de niveles")
async def ranking(interaction: discord.Interaction):
    data = cargar_datos()
    niveles = data.get("niveles", {})
    if not niveles:
        await interaction.response.send_message("Todavía no hay ranking.")
        return

    ordenado = sorted(niveles.items(), key=lambda x: x[1]["xp"], reverse=True)[:10]
    texto = ""
    for i, (uid, info) in enumerate(ordenado, 1):
        try:
            member = await interaction.guild.fetch_member(int(uid))
            m = member.mention
        except:
            m = info.get("nombre", "Usuario")
        texto += f"**#{i}** {m} — Nivel **{info['nivel']}** ({info['xp']} XP)\n"

    embed = discord.Embed(title="🏆 Ranking de Niveles", description=texto, color=discord.Color.gold())
    await interaction.response.send_message(embed=embed)

@bot.tree.command(name="poema", description="Te envía un poema aleatorio")
async def poema_cmd(interaction: discord.Interaction):
    await interaction.response.send_message(f"📜 **Poema para ti:**\n\n{random.choice(poemas)}")

@bot.tree.command(name="plantilla", description="Genera una plantilla de presentación")
@app_commands.describe(titulo="Título de la presentación")
async def plantilla(interaction: discord.Interaction, titulo: str = "Plantilla de Presentación"):
    await interaction.response.defer()
    try:
        archivo = crear_plantilla(titulo=titulo, autor=interaction.user.display_name)
        file = discord.File(archivo, filename=f"{titulo.replace(' ', '_')}.pptx")
        await interaction.followup.send(f"✅ Aquí tienes tu plantilla: **{titulo}**", file=file)
    except Exception as e:
        await interaction.followup.send(f"❌ Error: `{e}`")

def crear_plantilla(titulo: str = "Plantilla de Presentación", autor: str = "Discord Bot"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    azul_oscuro = RGBColor(15, 32, 64)
    azul = RGBColor(37, 99, 235)
    gris = RGBColor(100, 116, 139)
    blanco = RGBColor(255, 255, 255)

    def agregar_fondo(slide, color):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()

    def agregar_texto(slide, left, top, width, height, texto, size=24, bold=False, color=azul_oscuro, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = texto
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align

    # Portada
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    agregar_fondo(slide, azul_oscuro)
    barra = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.8), Inches(13.333), Inches(0.15))
    barra.fill.solid()
    barra.fill.fore_color.rgb = azul
    barra.line.fill.background()
    agregar_texto(slide, 0.8, 2.2, 11.5, 1, titulo, size=40, bold=True, color=blanco, align=PP_ALIGN.CENTER)
    agregar_texto(slide, 0.8, 3.2, 11.5, 0.6, f"Creado por {autor}", size=18, color=RGBColor(180, 200, 255), align=PP_ALIGN.CENTER)

    # Índice
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    agregar_fondo(slide, blanco)
    agregar_texto(slide, 0.8, 0.4, 11, 0.8, "Índice", size=32, bold=True, color=azul_oscuro)
    for i, tema in enumerate(["1. Introducción", "2. Objetivos", "3. Desarrollo", "4. Puntos clave", "5. Conclusiones", "6. Cierre"]):
        agregar_texto(slide, 1.2, 1.5 + i * 0.7, 10, 0.6, tema, size=22, color=gris)

    # Contenido
    for titulo_sec, contenido in [
        ("Introducción", "Escribe aquí la introducción.\n\n• Punto 1\n• Punto 2\n• Punto 3"),
        ("Objetivos", "• Objetivo 1\n• Objetivo 2\n• Objetivo 3"),
        ("Desarrollo", "Desarrolla aquí el contenido principal."),
        ("Puntos clave", "• Idea 1\n• Idea 2\n• Idea 3\n• Idea 4"),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        agregar_fondo(slide, blanco)
        barra = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        barra.fill.solid()
        barra.fill.fore_color.rgb = azul_oscuro
        barra.line.fill.background()
        agregar_texto(slide, 0.8, 0.3, 11, 0.7, titulo_sec, size=28, bold=True, color=blanco)
        agregar_texto(slide, 0.8, 1.6, 11.5, 4.5, contenido, size=20, color=gris)

    # Conclusión
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    agregar_fondo(slide, blanco)
    agregar_texto(slide, 0.8, 0.4, 11, 0.8, "Conclusiones", size=32, bold=True, color=azul_oscuro)
    agregar_texto(slide, 0.8, 1.6, 11.5, 4, "• Resumen\n• Mensaje final\n• Próximos pasos\n\n¡Gracias!", size=20, color=gris)

    # Cierre
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    agregar_fondo(slide, azul_oscuro)
    agregar_texto(slide, 0.8, 2.8, 11.5, 1, "¡Gracias!", size=44, bold=True, color=blanco, align=PP_ALIGN.CENTER)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer

@bot.event
async def on_ready():
    print(f"✅ Bot conectado como {bot.user}")
    try:
        synced = await bot.tree.sync()
        print(f"✅ Comandos sincronizados: {len(synced)}")
    except Exception as e:
        print(e)

    if not enviar_chistes.is_running():
        enviar_chistes.start()
    if not enviar_poemas.is_running():
        enviar_poemas.start()
    if not enviar_info.is_running():
        enviar_info.start()

if __name__ == "__main__":
    bot.run(TOKEN)
